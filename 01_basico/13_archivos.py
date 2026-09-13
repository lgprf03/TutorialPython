import csv
import json
import os
from openpyxl import Workbook, load_workbook

# 1. Archivos de Texto Plano (.txt)
def demo_txt() -> None:
    path = "notas.txt"

    # CREAR / ESCRIBIR ('w')
    with open(path, "w", encoding="utf-8") as file:
        file.write("Línea 1: Nota inicial\n")
        file.write("Línea 2: Segunda nota\n")

    # CARGAR / LEER ('r')
    with open(path, "r", encoding="utf-8") as file:
        contenido = file.read()
        print(f"--- TXT Inicial ---\n{contenido}")

    # MODIFICAR / AGREGAR ('a')
    with open(path, "a", encoding="utf-8") as file:
        file.write("Línea 3: Nota agregada posteriormente\n")

    # ELIMINAR
    if os.path.exists(path):
        os.remove(path)
        print("✓ Archivo .txt eliminado correctamente.\n")

# 2. Archivos JSON (.json)
def demo_json() -> None:
    path = "config.json"
    data_inicial = {"usuario": "admin", "theme": "dark", "intentos": 3}

    # CREAR / GUARDAR (json.dump)
    with open(path, "w", encoding="utf-8") as file:
        json.dump(data_inicial, file, indent=4)

    # CARGAR / LEER (json.load)
    with open(path, "r", encoding="utf-8") as file:
        config = json.load(file)
        print(f"--- JSON Cargado --- \nUsuario: {config['usuario']}")

    # MODIFICAR (Cargar -> Alterar dict -> Volver a guardar)
    config["theme"] = "light"
    config["intentos"] += 1

    with open(path, "w", encoding="utf-8") as file:
        json.dump(config, file, indent=4)

    # ELIMINAR
    if os.path.exists(path):
        os.remove(path)
        print("✓ Archivo .json eliminado correctamente.\n")

# 3. Archivos CSV (.csv)
def demo_csv() -> None:
    path = "usuarios.csv"
    filas = [
        ["id", "nombre", "rol"],
        [1, "Ana", "Developer"],
        [2, "Carlos", "Analyst"],
    ]

    # CREAR / ESCRIBIR (csv.writer)
    # newline='' evita líneas en blanco adicionales en Windows
    with open(path, "w", newline="", encoding="utf-8") as file:
        writer = csv.writer(file)
        writer.writerows(filas)

    # CARGAR / LEER (csv.reader)
    with open(path, "r", encoding="utf-8") as file:
        reader = csv.reader(file)
        print("--- CSV Contenido ---")
        for fila in reader:
            print(f"Fila: {fila}")

    # MODIFICAR (Agregar una fila con 'a')
    nueva_fila = [3, "Elena", "Designer"]
    with open(path, "a", newline="", encoding="utf-8") as file:
        writer = csv.writer(file)
        writer.writerow(nueva_fila)

    # ELIMINAR
    if os.path.exists(path):
        os.remove(path)
        print("✓ Archivo .csv eliminado correctamente.\n")

# 4. Archivo Excel (.xlsx)
def demo_xlsx() -> None:
    path = "reporte.xlsx"

    # CREAR
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"

    # Escribir encabezados y datos
    ws.append(["Producto", "Cantidad", "Precio"])
    ws.append(["Laptop", 2, 1200.00])
    ws.append(["Mouse", 10, 25.00])

    wb.save(path)
    wb.close()  # Es importante cerrar la instancia del libro

    # CARGAR Y LEER
    wb_cargado = load_workbook(path)
    ws_cargada = wb_cargado["Ventas"]

    print("--- Excel Contenido ---")
    for row in ws_cargada.iter_rows(values_only=True):
        print(row)

    # MODIFICAR (Cargar -> Editar celda o append -> Guardar)
    ws_cargada.append(["Teclado", 5, 45.00])
    ws_cargada["B2"] = 3  # Actualiza la cantidad de Laptops de 2 a 3

    wb_cargado.save(path)
    wb_cargado.close()

    # ELIMINAR
    if os.path.exists(path):
        os.remove(path)
        print("✓ Archivo .xlsx eliminado correctamente.\n")

if __name__ == "__main__":
    demo_txt()
    # demo_json()
    # demo_csv()
    # demo_xlsx()